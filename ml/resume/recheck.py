"""The tailored resume for one job, scored again from its own files: the job page's Re-check.

Three measures, none of which calls a model, so a re-check is free and takes seconds:

  * ATS keywords -- how many of the posting's keywords the text holds, the original resume and the
    tailored one side by side (ml/resume/ats.py);
  * requirements named word for word -- the posting's requirements the resume names outright, by
    the strict matching rules, again both resumes (ml/tailoring/rounds.requirements_shown);
  * ATS readability -- the PDF read back by pypdf, PDFium and pdfminer and the .docx by python-docx
    and its raw XML, every line checked whole and in order (ml/resume/ats_readback.py).

And what a reader would trip on in the text itself: a word broken by a space ("Cloud Bui ld"), a
skills row with nothing after its colon. Text copied out of a PDF comes back like that, and every
resume tailored from it inherits it.

The result is kept beside the tailored resume as <stem>_recheck.json.
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

HEADINGS = ("summary", "profile", "skills", "experience", "education", "certifications", "projects",
            "highlights", "expertise", "competencies", "awards", "publications")


def _files(cfg, job) -> tuple[Path, Path]:
    from ml.resume.packet import packet_root
    from ml.resume.writer import output_filename
    docx = packet_root(cfg) / output_filename(job.company, job.title, job.external_id or "")
    return docx.with_suffix(".pdf"), docx


def blocks_of(docx: Path) -> list[tuple[str, str]]:
    """What the page is meant to say, in order, read from the .docx the person sends."""
    from docx import Document
    from ml.resume.writer import _is_role_line
    lines = [p for p in Document(str(docx)).paragraphs if p.text.strip()]
    out: list[tuple[str, str]] = []
    for i, p in enumerate(lines):
        text = p.text.strip()
        if i == 0:
            out.append(("name", text))
        elif i == 1:
            out.append(("contact", text))
        elif p.style.name.lower().startswith("heading") or (text.isupper() and text.lower() in HEADINGS):
            out.append(("heading", text))
        elif _is_role_line(text) and re.search(r"(19|20)\d\d", text):
            out.append(("role", text))
        elif text.startswith(("•", "-", "*")):
            out.append(("bullet", text))
        else:
            out.append(("text", text))
    return out


_SPLIT_HYPHEN = re.compile(r"\b\w+ -\w+|\b\w+- \w+")
_COMMON = {"a", "an", "as", "at", "be", "by", "do", "go", "if", "in", "is", "it", "my", "no", "of", "on",
           "or", "so", "to", "up", "us", "we", "ai", "ui", "ux", "qa", "ci", "cd", "ml", "os", "db"}


def text_problems(docx: Path) -> list[str]:
    """Breaks a reader -- person or parser -- trips on, as the file says them."""
    from docx import Document
    found: list[str] = []
    for p in Document(str(docx)).paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if text.endswith(":") and len(text) < 120:
            found.append(f"“{text}” ends in a colon with nothing after it")
        for m in _SPLIT_HYPHEN.finditer(text):
            found.append(f"“{m.group(0)}” has a space inside a hyphenated word")
        # A short lower-case piece glued to a capitalised word: "Bui ld", "SonarQ ube", "Gra phDB".
        for m in re.finditer(r"\b([A-Z][a-zA-Z]{1,6}) ([a-z]{1,4}[A-Za-z]*)\b", text):
            tail = m.group(2)
            if tail.lower() not in _COMMON and len(m.group(1)) <= 6 and len(tail) <= 5 and not tail.startswith(("and", "the", "for", "with", "from")):
                joined = m.group(1) + tail
                if re.search(rf"\b{re.escape(joined)}\b", _VOCAB, re.I):
                    found.append(f"“{m.group(0)}” looks like one word split in two (“{joined}”)")
        for m in re.finditer(r"\b([a-z]{3,}) ([a-z]{2,4})\b", text):
            joined = m.group(1) + m.group(2)
            if m.group(2) not in _COMMON and re.search(rf"\b{re.escape(joined)}\b", _VOCAB, re.I) \
                    and not re.search(rf"\b{re.escape(m.group(2))}\b", _VOCAB, re.I):
                found.append(f"“{m.group(0)}” looks like one word split in two (“{joined}”)")
    return list(dict.fromkeys(found))


# Words a resume like this uses, to tell "Bui ld" (a break) from "Cloud Run" (two words).
_VOCAB = """build sonarqube graphdb scanning infrastructure kubernetes terraform ansible jenkins docker
deployment deployments pipeline pipelines monitoring observability security compliance automation
configuration provisioning containerization orchestration management engineering engineer platform
platforms reliability availability scalability performance optimization optimized integration
integrations architecture architectures governance vulnerability vulnerabilities migration migrations
development developer operations incident incidents recovery resilience networking database databases
analytics application applications environment environments cloudformation openshift opentelemetry
artifactory prometheus grafana datadog splunk cloudwatch gitlab github bitbucket kafka redis postgresql
""" + ""


def run(cfg, job, base_docx: Path | None = None) -> dict:
    """Score the job's tailored resume again; {} fields say what could not be read."""
    from ml.resume import ats, ats_readback
    from ml.resume.pipeline import base_resume_path, load_base_resume
    from ml.tailoring import requirements, rounds

    pdf, docx = _files(cfg, job)
    out: dict = {"checked_at": datetime.now().strftime("%Y-%m-%d %H:%M"), "has_resume": docx.exists() or pdf.exists()}
    if not out["has_resume"]:
        return out

    keywords = ats.check(cfg, job, load_base_resume(cfg).text())
    out["keywords"] = {"before": keywords.get("before"), "after": keywords.get("after"),
                       "missing": (keywords.get("fixable") or []) + (keywords.get("gaps") or [])[:8]}

    try:
        reqs, _ = requirements.extract(f"{job.title}\n\n{job.description or ''}", cfg, use_model=False)
        base = Path(base_docx or base_resume_path(cfg))
        out["requirements"] = {"count": len(reqs),
                               "before": round(rounds.requirements_shown(reqs, base), 1) if reqs else None,
                               "after": round(rounds.requirements_shown(reqs, docx), 1) if reqs and docx.exists() else None}
    except Exception as exc:
        out["requirements"] = {"error": f"{type(exc).__name__}: {str(exc)[:120]}"}

    if docx.exists():
        readings = ats_readback.read_back(pdf if pdf.exists() else None, docx, blocks_of(docx))
        out["readability"] = [{"reader": r.reader, "ok": not r.findings,
                               "problems": [f.detail for f in r.findings[:4]]} for r in readings]
        out["text_problems"] = text_problems(docx)[:12]

    pdf.with_name(pdf.stem + "_recheck.json").write_text(json.dumps(out, indent=1))
    return out


def last(cfg, job) -> dict | None:
    """The last check, unless the resume was written again since: then None, and the page checks afresh."""
    pdf, docx = _files(cfg, job)
    saved = pdf.with_name(pdf.stem + "_recheck.json")
    try:
        found = json.loads(saved.read_text())
        written = max((f.stat().st_mtime for f in (pdf, docx) if f.exists()), default=0)
        return None if written > saved.stat().st_mtime else found
    except (OSError, ValueError):
        return None
